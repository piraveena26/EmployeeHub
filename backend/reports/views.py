import csv
from datetime import timedelta
from decimal import Decimal
from django.http import HttpResponse
from django.utils import timezone
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from rest_framework import status
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from authentication.models import UserRole
from authentication.permissions import IsHRManager
from employees.models import Employee, EmploymentStatusChoices
from organization.models import Department
from attendance.models import AttendanceRecord, AttendanceStatus
from leave_management.models import LeaveRequest, LeaveStatus, LeaveBalance
from timesheets.models import Timesheet, TimesheetStatus
from payroll.models import Payslip, PayrollPeriod
from performance.models import PerformanceReview, Goal


class DashboardMetricsView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        user = request.user
        today = timezone.localdate()

        # Shared Employee view info if user is an employee
        employee = Employee.objects.select_related('department', 'designation', 'shift').filter(user=user).first()

        # Employee Dashboard Data
        employee_data = None
        if employee:
            today_att = AttendanceRecord.objects.filter(employee=employee, date=today).first()
            recent_payslip = Payslip.objects.filter(employee=employee).order_by('-created_at').first()
            balances = LeaveBalance.objects.filter(employee=employee, year=today.year).select_related('leave_type')
            pending_leaves = LeaveRequest.objects.filter(employee=employee, status=LeaveStatus.PENDING).count()
            active_goals = Goal.objects.filter(employee=employee, status__in=['NOT_STARTED', 'IN_PROGRESS']).count()

            employee_data = {
                'employee_id': employee.employee_id,
                'full_name': employee.full_name,
                'department': employee.department.name if employee.department else 'Unassigned',
                'designation': employee.designation.name if employee.designation else 'Staff',
                'photo': employee.photo_display,
                'today_checked_in': bool(today_att and today_att.check_in_time),
                'today_checked_out': bool(today_att and today_att.check_out_time),
                'today_check_in_time': str(today_att.check_in_time) if (today_att and today_att.check_in_time) else None,
                'today_hours': float(today_att.working_hours) if today_att else 0.0,
                'leave_balances': [
                    {
                        'type': b.leave_type.name,
                        'remaining': float(b.remaining_days),
                        'total': float(b.total_days),
                        'used': float(b.used_days)
                    } for b in balances
                ],
                'pending_leaves_count': pending_leaves,
                'active_goals_count': active_goals,
                'last_payslip_amount': float(recent_payslip.net_salary) if recent_payslip else 0.0,
                'last_payslip_number': recent_payslip.payslip_number if recent_payslip else None,
            }

        if user.is_regular_employee and not user.is_hr:
            return Response({
                'role': user.role,
                'employee_data': employee_data
            })

        # HR / Admin Dashboard Data
        total_employees = Employee.objects.count()
        active_employees = Employee.objects.filter(employment_status=EmploymentStatusChoices.ACTIVE).count()
        present_today = AttendanceRecord.objects.filter(
            date=today, status__in=[AttendanceStatus.PRESENT, AttendanceStatus.LATE]
        ).count()
        absent_today = max(0, active_employees - present_today)
        on_leave_today = LeaveRequest.objects.filter(
            status=LeaveStatus.APPROVED,
            start_date__lte=today,
            end_date__gte=today
        ).count()

        pending_leaves_count = LeaveRequest.objects.filter(status=LeaveStatus.PENDING).count()
        pending_timesheets_count = Timesheet.objects.filter(status=TimesheetStatus.SUBMITTED).count()

        latest_period = PayrollPeriod.objects.order_by('-year', '-month').first()
        payroll_status = {
            'period': f"{latest_period.month:02d}/{latest_period.year}" if latest_period else "None",
            'is_processed': latest_period.is_processed if latest_period else False,
            'total_amount': sum([float(p.net_salary) for p in latest_period.payslips.all()]) if (latest_period and latest_period.payslips.exists()) else 0.0
        }

        # Department distribution chart
        dept_distribution = []
        for d in Department.objects.all():
            cnt = d.employees.filter(employment_status=EmploymentStatusChoices.ACTIVE).count()
            if cnt > 0:
                dept_distribution.append({'name': d.name, 'count': cnt})

        # Attendance weekly trend (past 7 days)
        attendance_trend = []
        for i in range(6, -1, -1):
            day = today - timedelta(days=i)
            present_cnt = AttendanceRecord.objects.filter(
                date=day, status__in=[AttendanceStatus.PRESENT, AttendanceStatus.LATE]
            ).count()
            attendance_trend.append({
                'day': day.strftime('%a, %b %d'),
                'present': present_cnt,
                'total': active_employees
            })

        return Response({
            'role': user.role,
            'metrics': {
                'total_employees': total_employees,
                'active_employees': active_employees,
                'present_today': present_today,
                'absent_today': absent_today,
                'on_leave_today': on_leave_today,
                'pending_leaves_count': pending_leaves_count,
                'pending_timesheets_count': pending_timesheets_count,
                'payroll_status': payroll_status,
                'department_distribution': dept_distribution,
                'attendance_trend': attendance_trend,
            },
            'employee_data': employee_data
        })


class ExportEmployeesView(APIView):
    permission_classes = [IsHRManager]

    def get(self, request):
        fmt = request.query_params.get('format', 'csv')
        employees = Employee.objects.all().select_related('department', 'designation', 'shift')

        if fmt == 'excel':
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Employees"
            headers = ["Employee ID", "Full Name", "Email", "Phone", "Department", "Designation", "Employment Type", "Status", "Joining Date", "Salary"]
            ws.append(headers)
            header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
            for cell in ws[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = header_fill

            for emp in employees:
                ws.append([
                    emp.employee_id, emp.full_name, emp.email, emp.phone or '',
                    emp.department.name if emp.department else '',
                    emp.designation.name if emp.designation else '',
                    emp.get_employment_type_display(),
                    emp.get_employment_status_display(),
                    str(emp.joining_date),
                    float(emp.basic_salary)
                ])

            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = 'attachment; filename="EmployeeHub_Employees.xlsx"'
            wb.save(response)
            return response
        else:
            response = HttpResponse(content_type='text/csv')
            response['Content-Disposition'] = 'attachment; filename="EmployeeHub_Employees.csv"'
            writer = csv.writer(response)
            writer.writerow(["Employee ID", "Full Name", "Email", "Phone", "Department", "Designation", "Employment Type", "Status", "Joining Date", "Salary"])
            for emp in employees:
                writer.writerow([
                    emp.employee_id, emp.full_name, emp.email, emp.phone or '',
                    emp.department.name if emp.department else '',
                    emp.designation.name if emp.designation else '',
                    emp.get_employment_type_display(),
                    emp.get_employment_status_display(),
                    str(emp.joining_date),
                    emp.basic_salary
                ])
            return response


class ExportAttendanceView(APIView):
    permission_classes = [IsHRManager]

    def get(self, request):
        records = AttendanceRecord.objects.all().select_related('employee', 'employee__department')
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="EmployeeHub_Attendance.csv"'
        writer = csv.writer(response)
        writer.writerow(["Date", "Employee ID", "Name", "Department", "Check In", "Check Out", "Hours", "Status", "Late"])
        for r in records:
            writer.writerow([
                str(r.date), r.employee.employee_id, r.employee.full_name,
                r.employee.department.name if r.employee.department else '',
                str(r.check_in_time or ''), str(r.check_out_time or ''),
                r.working_hours, r.get_status_display(), "Yes" if r.is_late else "No"
            ])
        return response
